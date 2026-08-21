#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown → DOCX 转换器（中文优化）
==================================

把仓库里的指南类 Markdown 转成排版整齐的 Word 文档，让 .md 保持唯一真源。

支持：# 标题（1-4 级）、段落、- / * 无序列表、1. 有序列表、
表格（GFM 管道语法）、``` 代码块、> 引用块、--- 分隔线、
行内 **粗体** `代码` 与链接文本。

用法
----
    python scripts/md2docx.py docs/mcp-agv/gmx/GROMACS-MCP指南.md
    python scripts/md2docx.py docs/mcp-agv/**/*.md          # 批量
    python scripts/md2docx.py input.md -o output.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CJK_FONT = "微软雅黑"
LATIN_FONT = "Calibri"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x14, 0x4E, 0x5A)


def _style_run(run, *, mono=False, size=10.5, bold=False, italic=False, color=None):
    run.font.name = MONO_FONT if mono else LATIN_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), MONO_FONT if mono else CJK_FONT)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")


def _add_inline(par, text: str, base_size=10.5):
    """处理行内 **粗体**、`代码`、[文本](链接)、*斜体*。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _style_run(par.add_run(part[2:-2]), size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            _style_run(r, mono=True, size=base_size - 1)
            r.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        elif part.startswith("[") and "](" in part:
            label = part[1: part.index("]")]
            url = part[part.index("](") + 2: -1]
            r = par.add_run(label)
            _style_run(r, size=base_size, color=RGBColor(0x1A, 0x5F, 0xB4))
            r.underline = True
            if url not in label:
                _style_run(par.add_run(f" ({url})"), size=base_size - 1.5,
                           color=RGBColor(0x88, 0x88, 0x88))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _style_run(par.add_run(part[1:-1]), size=base_size, italic=True)
        else:
            _style_run(par.add_run(part), size=base_size)


def _shade(par, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    par._p.get_or_add_pPr().append(shd)


def _init_doc() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = LATIN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25
    for lvl, size in [(1, 19), (2, 14.5), (3, 12), (4, 11)]:
        s = doc.styles[f"Heading {lvl}"]
        s.font.name = LATIN_FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = ACCENT
        s.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    return doc


def _add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            cell = cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            txt = row[j] if j < len(row) else ""
            txt = txt.replace("<br>", " ")
            if i == 0:
                _style_run(par.add_run(txt), size=9.5, bold=True)
            else:
                _add_inline(par, txt, base_size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path: Path, out_path: Path) -> tuple[int, int, int]:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = _init_doc()

    n_head = n_table = n_code = 0
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush_table():
        nonlocal table_buf, n_table
        if table_buf:
            _add_table(doc, table_buf)
            n_table += 1
            table_buf = []

    while i < len(lines):
        line = lines[i]

        # 代码块
        if line.lstrip().startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Inches(0.18)
                par.paragraph_format.space_before = Pt(4)
                par.paragraph_format.space_after = Pt(8)
                par.paragraph_format.line_spacing = 1.0
                _style_run(par.add_run("\n".join(code_buf)), mono=True, size=8.8)
                _shade(par, "F2F4F7")
                code_buf = []
                in_code = False
                n_code += 1
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue  # 分隔行
            table_buf.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            par = doc.add_paragraph()
            _style_run(par.add_run("─" * 46), size=8,
                       color=RGBColor(0xBB, 0xBB, 0xBB))
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).replace("**", "").replace("`", "")
            if level == 1 and n_head == 0:
                par = doc.add_paragraph()
                _style_run(par.add_run(text), size=22, bold=True, color=ACCENT)
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(10)
            else:
                doc.add_heading(text, level=min(level, 4))
            n_head += 1
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            body = [stripped.lstrip("> ").rstrip()]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith(">"):
                i += 1
                body.append(lines[i].strip().lstrip("> ").rstrip())
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.22)
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(8)
            _add_inline(par, " ".join(x for x in body if x), base_size=10)
            _shade(par, "E8F4F2")
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+])\s+(.*)", line)
        if m:
            indent = len(m.group(1))
            style = "List Bullet" if indent < 2 else "List Bullet 2"
            par = doc.add_paragraph(style=style)
            par.paragraph_format.space_after = Pt(2)
            _add_inline(par, m.group(3))
            i += 1
            continue

        m = re.match(r"^(\s*)\d+\.\s+(.*)", line)
        if m:
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(2)
            _add_inline(par, m.group(2))
            i += 1
            continue

        # 普通段落
        par = doc.add_paragraph()
        _add_inline(par, stripped)
        i += 1

    flush_table()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return n_head, n_table, n_code


def main() -> int:
    ap = argparse.ArgumentParser(description="Markdown → DOCX（中文优化）")
    ap.add_argument("inputs", nargs="+", help="markdown 文件路径")
    ap.add_argument("-o", "--out", help="输出路径（仅单文件时有效）")
    args = ap.parse_args()

    if args.out and len(args.inputs) > 1:
        print("-o 只能配合单个输入文件使用", file=sys.stderr)
        return 2

    for raw in args.inputs:
        src = Path(raw)
        if not src.exists():
            print(f"找不到文件：{src}", file=sys.stderr)
            return 1
        dst = Path(args.out) if args.out else src.with_suffix(".docx")
        h, t, c = convert(src, dst)
        print(f"{src}  →  {dst}   （标题 {h}，表格 {t}，代码块 {c}）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

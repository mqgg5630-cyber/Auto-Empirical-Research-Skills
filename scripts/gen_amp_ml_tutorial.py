# -*- coding: utf-8 -*-
"""生成《抗菌肽机器学习预测 · MCP + Skills 全流程实操教学》docx。"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "抗菌肽机器学习预测-MCP与Skills实操教学.docx"
doc = Document()

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.25
for i, sz in [(1, 19), (2, 14.5), (3, 12)]:
    s = doc.styles[f"Heading {i}"]
    s.font.name = "Calibri"
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0x14, 0x4E, 0x5A)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def p(t="", bold=False, italic=False, size=10.5, color=None, align=None, after=6):
    par = doc.add_paragraph()
    r = par.add_run(t)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        r.font.color.rgb = color
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    return par


def b(t):
    par = doc.add_paragraph(style="List Bullet")
    r = par.add_run(t)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    par.paragraph_format.space_after = Pt(2)
    return par


def num(t):
    par = doc.add_paragraph(style="List Number")
    r = par.add_run(t)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    par.paragraph_format.space_after = Pt(2)
    return par


def code(t):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.18)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.line_spacing = 1.0
    r = par.add_run(t)
    r.font.name = "Consolas"
    r.font.size = Pt(8.8)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F4F7")
    par._p.get_or_add_pPr().append(shd)
    return par


def prompt(t):
    """给用户直接粘贴进反重力的指令块。"""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.18)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.line_spacing = 1.15
    r = par.add_run("▶ 指令：" + t)
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.italic = True
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.color.rgb = RGBColor(0x0B, 0x3D, 0x4A)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "E8F4F2")
    par._p.get_or_add_pPr().append(shd)
    return par


def table(headers, rows, widths=None, fs=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(fs + 0.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.name = "Calibri"
            r.font.size = Pt(fs)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ============ 封面 ============
p("抗菌肽（AMP）机器学习预测", bold=True, size=24, color=RGBColor(0x14, 0x4E, 0x5A),
  align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("MCP + Agent Skills 全流程实操教学", bold=True, size=15,
  color=RGBColor(0x2A, 0x6F, 0x7A), align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
p("从数据库检索 → 防泄漏清洗 → 特征工程 → 建模调参 → 可解释性 → 论文与复现包",
  size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))
p("运行环境：Google Antigravity IDE ｜ 技能库：Auto-Empirical-Research-Skills",
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x77, 0x77, 0x77))
p("2026-08-20", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x88, 0x88, 0x88), after=12)

p("怎么用这份文档：每一章都是「先讲清楚这一步在干什么和为什么、要用哪些工具、然后给一段可以直接粘进反重力的指令、"
  "最后给验收标准」。你不需要一次读完 —— 按章节顺序做，做完一章再看下一章。",
  size=10, italic=True, color=RGBColor(0x44, 0x44, 0x44))

# ============ 0 环境 ============
doc.add_heading("第 0 章　环境准备（30 分钟一次性配置）", level=1)

doc.add_heading("0.1 装什么", level=2)
table(
    ["工具", "在本课题的角色", "费用"],
    [
        ["Antigravity IDE", "主战场，Agent 可并行跑多任务", "免费"],
        ["Hugging Face MCP", "找 ESM-2 / ProtT5 等蛋白语言模型与肽段数据集", "免费（注册号取 token）"],
        ["Optuna MCP", "超参搜索 + 参数重要性图 / 优化历史图", "免费"],
        ["Jupyter MCP", "真跑训练 cell，读回指标与图", "免费"],
        ["Context7", "取 torch / transformers / scikit-learn / xgboost 当前版本 API", "免费"],
        ["paper-search + Zotero", "AMP 预测文献检索与引用可信化", "免费"],
        ["Markitdown", "把论文 PDF、数据库说明书转成可读文本", "免费"],
        ["DBHub（只读）", "把清洗后的肽段库放 SQLite，随时查子集", "免费"],
        ["filesystem + memory", "整理 FASTA 文件、记录每次实验决策", "免费"],
        ["W&B MCP（可选）", "多次训练的实验追踪与对比", "学术免费"],
    ],
    widths=[1.5, 4.0, 1.5],
)

doc.add_heading("0.2 mcp_config.json（粘到 ~/.gemini/config/mcp_config.json）", level=2)
code('''{
  "mcpServers": {
    "huggingface": {
      "serverUrl": "https://huggingface.co/mcp",
      "headers": { "Authorization": "Bearer ${HF_TOKEN}" }
    },
    "optuna": {
      "command": "uvx",
      "args": ["optuna-mcp", "--storage", "sqlite:////ABS/PATH/amp-project/optuna.db"]
    },
    "jupyter": {
      "command": "uvx",
      "args": ["jupyter-mcp-server"],
      "env": { "JUPYTER_URL": "http://localhost:8888", "JUPYTER_TOKEN": "${JUPYTER_TOKEN}" }
    },
    "context7":   { "command": "npx", "args": ["-y", "@upstash/context7-mcp"] },
    "markitdown": { "command": "uvx", "args": ["markitdown-mcp"] },
    "paper-search": {
      "command": "uvx",
      "args": ["--from", "git+https://github.com/openags/paper-search-mcp", "paper-search-mcp"],
      "env": { "PAPER_SEARCH_MCP_UNPAYWALL_EMAIL": "your@email.edu" }
    },
    "zotero": {
      "command": "uvx",
      "args": ["--from", "git+https://github.com/54yyyu/zotero-mcp", "zotero-mcp"],
      "env": { "ZOTERO_LOCAL": "true" }
    },
    "dbhub": {
      "command": "npx",
      "args": ["-y", "@bytebase/dbhub", "--transport", "stdio",
               "--dsn", "sqlite:////ABS/PATH/amp-project/data/amp.db", "--readonly"]
    },
    "filesystem": {
      "command": "npx",
      "args": ["-y", "@modelcontextprotocol/server-filesystem", "/ABS/PATH/amp-project"]
    },
    "memory": {
      "command": "uvx", "args": ["basic-memory", "mcp"],
      "env": { "BASIC_MEMORY_HOME": "/ABS/PATH/amp-project/notes" }
    }
  }
}''')
p("把 /ABS/PATH 换成真实绝对路径；HF_TOKEN 等先在 ~/.zshrc 里 export 好再启动反重力。"
  "改完到 Manage MCP Servers 点 Refresh。", size=10)

doc.add_heading("0.3 装 Skills（反重力 v1.14.2+ 原生支持 SKILL.md）", level=2)
code('''cd /ABS/PATH/amp-project
mkdir -p .agents/skills

# 从本仓库拷贝本课题最相关的技能包
REPO=/path/to/Auto-Empirical-Research-Skills
cp -r $REPO/skills/00.1-Full-empirical-analysis-skill_Python  .agents/skills/
cp -r $REPO/skills/03-K-Dense-AI-claude-scientific-skills/*   .agents/skills/
cp -r $REPO/skills/36-taoyunudt-literature-review-skill       .agents/skills/
cp -r $REPO/skills/52-keemanxp-slr-prisma                     .agents/skills/
cp -r $REPO/skills/54-scdenney-open-science-skills            .agents/skills/
cp -r $REPO/skills/72-kaggle-research                         .agents/skills/

# 另可启用 Google 官方 DeepMind Science Skills（38 个，接 AlphaFold/UniProt/PubMed/PDB）
# Settings → Customizations → Build with Google Plugins → Science''')

table(
    ["Skill", "在本课题干什么"],
    [
        ["00.1-Full-empirical-analysis_Python", "references/05-modeling.md 建模规范、06-robustness.md 稳健性清单、08-tables-plots.md 出图表"],
        ["03-K-Dense scientific skills", "hypothesis-generation 帮你把「预测抗菌肽」细化成可检验的假设；scientific-writing 出稿"],
        ["36 / 52（综述 + PRISMA）", "系统梳理已有 AMP 预测器，形成可复现的检索式与筛选流程"],
        ["54-open-science-skills", "数据可得性声明、复现包规范、预注册"],
        ["72-kaggle-research", "通过官方 Kaggle CLI 受控检索下载外部肽段数据集，带审计留痕"],
        ["DeepMind Science（官方）", "UniProt / PDB / PubMed 直连，构造负样本和查结构时用得上"],
    ],
    widths=[2.3, 4.7],
)
p("分工原则：Skill 管「按什么规范做」，MCP 管「能碰到什么外部资源」。Skill 会在步骤里主动调 MCP 工具。", size=10, italic=True)

doc.add_page_break()

# ============ 1 研究设计 ============
doc.add_heading("第 1 章　把课题拆成可执行的研究设计", level=1)

doc.add_heading("1.1 先明确你到底要预测什么", level=2)
p("“机器学习预测抗菌肽”至少有四种完全不同的任务，先选定一个，否则后面全乱：", after=4)
table(
    ["任务类型", "标签", "难度", "评价指标"],
    [
        ["A. AMP / non-AMP 二分类", "0/1", "入门，文献最多，容易做但也最卷", "MCC、AUC、AUPRC、SEN/SPE"],
        ["B. 活性强度回归（MIC 预测）", "log MIC 连续值", "中等，数据少但更有价值", "PCC、R²、RMSE、MAE"],
        ["C. 菌种特异性预测", "对某菌株是否有效", "中高，需 species-aware 数据", "分菌株分别报 MCC"],
        ["D. 功能亚类 / 毒性联合预测", "多标签", "高，需溶血性、细胞毒性标注", "macro-F1、每标签 AUC"],
    ],
    widths=[1.9, 1.7, 1.9, 1.5],
)
p("给你的建议：先做 A 建立完整流水线（两周内可跑通），再把同一套流水线迁移到 B。"
  "B（MIC 回归）在审稿人眼里比又一个二分类器更有增量价值——因为 A 类文献已经非常拥挤，"
  "近年报告的准确率普遍在 0.87–0.96 区间，很难再靠指标取胜。", after=4)

doc.add_heading("1.2 这个领域最大的坑：负样本选择偏差", level=2)
p("必须知道的一篇文章：Sidorczuk 等人在 Briefings in Bioinformatics (2022) 系统检验了 AMP 预测的基准数据，"
  "结论是——现有模型的性能高度依赖负样本是怎么采的。当训练集与测试集用同一种负样本采样策略时性能虚高，"
  "换一种采样策略性能就掉。作者的原话结论是：已发表的 AMP 预测基准比较都是不公平的，"
  "我们其实并不知道哪个模型最准。", after=4)
p("这意味着两件事：", bold=True, after=2)
b("① 你的论文必须明确交代负样本构造规则（关键词排除、长度分布匹配、亚细胞定位过滤、CD-HIT 阈值），并且最好用两套不同策略的负样本各评一遍。")
b("② 这本身就是一个可发表的切入点：如果你能提出更稳健的负样本构造方案或跨采样策略的评测框架，比再刷 0.5% 准确率有意义得多。")

prompt("用 paper-search 检索 2022 年以来关于 antimicrobial peptide prediction 的负样本构造与基准偏差问题的文献，"
       "以及 2024–2026 年新发表的 AMP 预测器（覆盖 arXiv、PubMed、Europe PMC、Crossref）。"
       "对每篇抽取：负样本来源、长度过滤区间、CD-HIT 阈值、是否做独立测试集、报告的 MCC。整理成对比表，"
       "标出哪些研究没有交代负样本策略。结果写进 Zotero 并在 memory 里记一条检索日志。")

doc.add_heading("1.3 三条候选研究问题（选一个作为主线）", level=2)
table(
    ["方向", "研究问题", "创新点", "可行性"],
    [
        ["稳健性", "跨负样本采样策略时，哪类模型性能最稳定？", "提出跨策略评测协议", "高，只需复现现有模型"],
        ["表征", "ESM-2 嵌入相比手工特征（AAC/CTD/理化）的增量究竟有多大，在小样本下是否仍成立？", "系统消融 + 样本量曲线", "高"],
        ["应用", "菌种特异性 MIC 回归 + 可解释性，定位决定活性的关键残基模式", "从「是不是 AMP」推进到「对谁多有效、为什么」", "中，需 GRAMPA 类数据"],
    ],
    widths=[1.0, 3.0, 1.8, 1.2],
)

doc.add_page_break()

# ============ 2 数据 ============
doc.add_heading("第 2 章　数据获取", level=1)

doc.add_heading("2.1 正样本：主流 AMP 数据库", level=2)
table(
    ["数据库", "规模量级", "特点", "备注"],
    [
        ["DBAASP v3", "约 2.2 万条", "带活性数据（MIC）、靶标菌株、实验条件，做回归首选", "dbaasp.org"],
        ["dbAMP", "约 3.5 万条", "整合面最广，常作参考库", "含功能注释"],
        ["DRAMP 3.0", "约 3 万条", "分天然/合成/专利，结构信息较全", "—"],
        ["APD3", "约 5 千条", "人工审编质量高，规模小", "经典正样本源"],
        ["CAMPR4", "约 2.4 万条", "含预测工具与序列库", "—"],
        ["LAMP2 / YADAMP", "各数千至万级", "常与上述合并去重", "—"],
        ["GRAMPA", "数千条 MIC 记录", "MIC 回归任务的常用公开集", "GitHub 可下载"],
    ],
    widths=[1.3, 1.2, 2.9, 1.6],
)
p("注意：这些库彼此高度重叠（同一条肽在多个库里重复出现），合并后必须去重，否则会把重复序列同时放进训练集和测试集，"
  "造成严重的性能虚高。", size=10)

doc.add_heading("2.2 负样本：怎么构造才站得住", level=2)
p("通用做法是从 UniProt / SwissProt 抽取非 AMP 蛋白片段，但细节决定成败。文献里常见的规则组合：", after=4)
b("关键词排除：剔除注释含 antimicrobial、antibiotic、antiviral、fungicide、defensin、bacteriocin、secreted 等的条目。")
b("长度分布匹配：让负样本长度分布与正样本一致（否则模型可能只是在学「长度」这个捷径特征）。")
b("亚细胞定位过滤：部分研究只保留胞质蛋白，以降低误标为 AMP 的概率。")
b("冗余去除：CD-HIT，常见阈值 0.4（严格）或 0.9（宽松）。阈值越严，任务越难，但结果越可信。")
b("类别平衡：正负 1:1 便于比较，但真实场景是极度不平衡的——建议主结果用平衡集，附录补一个 1:10 的不平衡评测。")

prompt("在 jupyter 里执行：从 UniProt 下载 SwissProt 全库，按上述五条规则构造负样本集。"
       "关键词排除清单写成配置项方便日后修改；长度分布用与正样本相同的直方图做分层抽样；"
       "分别生成 CD-HIT 0.4 与 0.9 两个版本。输出每一步的样本量流失表（类似 PRISMA 流程图），保存为 CSV。")

doc.add_heading("2.3 目录结构（先建好，后面所有脚本都依赖它）", level=2)
code('''amp-project/
├── .agents/skills/            # 装进来的 SKILL.md 技能包
├── data/
│   ├── raw/                   # 各数据库原始下载（只读，永不修改）
│   │   ├── dbaasp_v3.json
│   │   ├── dramp3.fasta
│   │   └── swissprot.fasta
│   ├── interim/               # 中间产物（去重、过滤后）
│   ├── processed/             # 最终训练用数据集
│   │   ├── pos_cdhit04.fasta
│   │   ├── neg_cdhit04.fasta
│   │   └── splits/            # train/valid/test 索引，固定随机种子
│   └── amp.db                 # SQLite，供 DBHub 查询
├── notebooks/                 # 探索性分析
├── src/
│   ├── data/                  # 下载与清洗脚本
│   ├── features/              # 特征提取
│   ├── models/                # 训练与评估
│   └── viz/                   # 出图
├── results/
│   ├── figures/
│   ├── tables/
│   └── metrics/               # 每次实验的 JSON 指标
├── notes/                     # basic-memory 研究日志
├── environment.yml
└── README.md''')

prompt("用 filesystem 按上面的结构建好目录，生成 .gitignore（忽略 data/raw、data/interim、*.pt、wandb/、__pycache__），"
       "并写一份 README 说明每个目录的用途和脚本运行顺序。")

doc.add_page_break()

# ============ 3 清洗 ============
doc.add_heading("第 3 章　清洗与防数据泄漏（最关键的一章）", level=1)
p("这一章做不好，后面所有指标都是假的。审稿人最常挑的也正是这里。", bold=True, color=RGBColor(0xA0, 0x30, 0x30))

doc.add_heading("3.1 标准清洗流程", level=2)
num("合并多库正样本 → 按序列完全一致去重（注意大小写与空白字符）。")
num("长度过滤：常见区间 5–100 aa（做二分类）或按你的任务调整。记录被剔除的数量。")
num("字符过滤：剔除含 B/J/O/U/X/Z 等非标准氨基酸的序列（或单独作为一个子集分析）。")
num("同源冗余去除：CD-HIT 在正样本内部、负样本内部各跑一次。")
num("跨集去冗余：正负样本之间也要检查，避免同一段序列同时出现在两类里。")
num("划分数据集：这一步是重点，见 3.2。")

doc.add_heading("3.2 划分数据集：随机划分是错的", level=2)
p("如果你随机 8:2 划分，训练集里某条肽的近似同源序列很可能落在测试集里，模型只要「记住」就能答对，"
  "测试指标会显著虚高。正确做法是同源性感知划分：", after=4)
b("先用 CD-HIT 把所有序列聚类（如 40% 相似度阈值），然后按「簇」而不是按「序列」划分训练/验证/测试。")
b("同一个簇的所有序列必须落在同一个 split 里。")
b("另外保留一个真正的独立测试集：最好来自不同数据库或不同发表年份（例如用 2024 年后新收录的肽做外部验证），这是最有说服力的。")

prompt("用 context7 查 CD-HIT 与 scikit-learn GroupKFold 的当前用法，然后在 jupyter 里实现同源性感知划分："
       "先 cd-hit 聚类（阈值 0.4）得到簇标签，再用 GroupShuffleSplit 按簇划分 train/valid/test = 7:1.5:1.5，"
       "固定随机种子 42。输出：各 split 的样本量、长度分布对比图、跨 split 的最大序列相似度（应远低于阈值）。"
       "把划分索引存进 data/processed/splits/，并在 memory 里记录本次划分的全部参数。")

doc.add_heading("3.3 验收标准（做完自查）", level=2)
table(
    ["检查项", "通过标准"],
    [
        ["跨 split 序列相似度", "任意 train-test 序列对的相似度 < CD-HIT 阈值"],
        ["重复序列", "全库唯一序列数 = 总条数"],
        ["长度分布", "正负样本、各 split 之间分布无显著差异（KS 检验 p > 0.05）"],
        ["氨基酸组成", "各 split 的 20 种氨基酸频率差异 < 2%"],
        ["样本流失记录", "有一张从原始到最终的完整流失表，每一步都能解释"],
        ["随机种子", "全流程固定，重跑结果完全一致"],
    ],
    widths=[2.6, 4.4],
)

doc.add_page_break()

# ============ 4 特征 ============
doc.add_heading("第 4 章　特征工程", level=1)

doc.add_heading("4.1 三层特征体系", level=2)
table(
    ["层级", "特征", "维度量级", "说明"],
    [
        ["手工-组成", "AAC 氨基酸组成、DPC 二肽组成、伪氨基酸组成 PseAAC", "20 / 400 / 20+λ", "最基础，可解释性强"],
        ["手工-理化", "净电荷、疏水矩、两亲性、等电点、Boman 指数、螺旋倾向", "10–50", "抗菌机制上有明确生物学意义，写讨论时好用"],
        ["手工-结构化", "CTD（组成-转换-分布）、PSSM 衍生特征", "100–500", "较老但仍常见于对比基线"],
        ["语言模型嵌入", "ESM-2、ProtT5、ProtBERT 的序列级向量", "320–2560", "近年主流，通常带来最大提升"],
    ],
    widths=[1.2, 2.7, 1.1, 2.0],
)
p("经验规律（来自近年多篇对比研究）：蛋白语言模型嵌入 > 手工特征；但两者拼接往往还能再涨一点，"
  "而且手工特征让你在讨论部分能讲出生物学故事（比如净电荷与膜穿透的关系）。所以别只用嵌入。", size=10)

prompt("用 huggingface 搜索适合短肽（5–100 aa）的蛋白语言模型，比较 ESM-2 的各个尺寸（8M/35M/150M/650M）"
       "在参数量、嵌入维度、显存需求上的差异，给出在单张 16GB 显卡上的推荐选择和理由。")

prompt("在 jupyter 里实现特征提取模块：(1) AAC + DPC + 理化性质（用 peptides 或 propy 库）；"
       "(2) ESM-2 嵌入（mean pooling 与 CLS 两种池化方式都实现）。"
       "每种特征保存为独立的 .npy 并记录 shape 与提取耗时。用 context7 确认 transformers 当前版本加载 ESM-2 的正确写法。")

doc.add_heading("4.2 一个容易被忽略的细节", level=2)
p("特征标准化必须只在训练集上 fit，然后 transform 到验证/测试集。用 sklearn Pipeline 把标准化和模型包在一起，"
  "再放进交叉验证——否则标准化过程已经偷看了测试集的分布，属于典型的隐性泄漏。", after=4)

doc.add_page_break()

# ============ 5 建模 ============
doc.add_heading("第 5 章　建模与调参", level=1)

doc.add_heading("5.1 从简到繁的四级阶梯", level=2)
table(
    ["级别", "模型", "特征", "作用"],
    [
        ["L0 基线", "逻辑回归、朴素贝叶斯", "AAC", "给出性能下界，如果复杂模型只比它好一点，说明有问题"],
        ["L1 传统 ML", "随机森林、XGBoost、SVM（RBF）", "手工特征拼接", "强基线，很多论文的最终模型就在这一级"],
        ["L2 嵌入 + ML", "RF / XGBoost / MLP", "ESM-2 嵌入", "性价比最高，不需要微调大模型"],
        ["L3 深度模型", "CNN、BiLSTM、微调 ESM-2", "原始序列", "上限最高，但要显卡、要防过拟合"],
    ],
    widths=[0.9, 2.1, 1.7, 2.3],
)
p("务必按顺序做，并且每一级都记录结果。论文里的消融表就是这么来的。跳过 L0/L1 直接上 L3 是新手最常见的错误——"
  "你会失去解释「复杂模型到底带来了多少增量」的能力。", size=10)

doc.add_heading("5.2 用 Optuna MCP 调参", level=2)
prompt("创建一个 Optuna study 名为 amp-xgb-cdhit04，方向为 maximize，目标是验证集 MCC。"
       "搜索空间：max_depth 整数 [3,10]、learning_rate 对数尺度 [0.01,0.3]、subsample [0.6,1.0]、"
       "colsample_bytree [0.6,1.0]、n_estimators 整数 [200,1500]、scale_pos_weight [1,5]。"
       "跑 60 个 trial：每个 trial 由你 suggest 参数，我在 jupyter 里训练并把 MCC 回报给你。"
       "跑完输出优化历史图、参数重要性图、平行坐标图，并告诉我哪两个参数最关键。")

p("Optuna MCP 的价值在于：它把「建 study → 建议参数 → 回报结果 → 可视化」都变成了对话，"
  "你不用写 objective 函数的样板代码，而且中途可以随时问它「目前最好的组合是什么」「参数重要性怎么变了」。", size=10)

doc.add_heading("5.3 防过拟合的硬性要求", level=2)
b("嵌套交叉验证：外层评估性能，内层调参。只用一层 CV 调参再报告同一层的结果，是乐观偏差的来源。")
b("早停：深度模型必须用验证集早停，并记录实际停在第几个 epoch。")
b("多次随机种子：至少跑 5 个不同种子，报告均值 ± 标准差。单次结果不可信。")
b("测试集只用一次：调参阶段绝对不碰测试集。反复看测试集调整策略，等于把测试集变成了验证集。")

doc.add_page_break()

# ============ 6 评估 ============
doc.add_heading("第 6 章　评估与结果呈现", level=1)

doc.add_heading("6.1 指标选择", level=2)
table(
    ["指标", "为什么要报", "注意"],
    [
        ["MCC", "类别不平衡下最稳健的单一指标，AMP 领域的事实标准", "主指标，摘要里必须有"],
        ["AUC (ROC)", "阈值无关的排序能力", "不平衡时会偏乐观"],
        ["AUPRC", "不平衡场景下比 AUC 更如实", "1:10 评测时必报"],
        ["SEN / SPE", "分别反映漏检与误报", "很多模型 SEN 很高但 SPE 只有 0.5 左右，一定要分开看"],
        ["F1 / Accuracy", "便于与文献比较", "单独看会误导，不要只报这两个"],
    ],
    widths=[1.0, 3.4, 2.6],
)
p("特别提醒：文献中不少 AMP 预测器的敏感度（SEN）高达 0.95 以上，但特异度（SPE）只有 0.5 左右——"
  "也就是把大量非 AMP 也判成了 AMP。如果你的模型也是这样，必须在讨论里坦诚说明，而不是只报总体准确率。", size=10)

doc.add_heading("6.2 必做的三组对比", level=2)
num("内部对比（消融）：L0 → L1 → L2 → L3 逐级提升多少；手工特征 vs 嵌入 vs 拼接。")
num("外部对比（与已有工具）：至少与 AMPScanner V2、amPEPpy、AmpGram 等公开可跑的工具在同一测试集上比较。"
    "注意要用它们的原始模型跑你的测试集，而不是引用它们论文里的数字——那是在不同数据上得到的。")
num("稳健性对比：换一套负样本采样策略、换 CD-HIT 阈值（0.4 vs 0.9），看排名是否变化。这正是第 1.2 节说的那个坑。")

prompt("加载 results/metrics/ 下所有实验的 JSON，生成三张表："
       "(1) 消融表：模型 × 特征组合 × 各指标（均值±标准差，5 个种子）；"
       "(2) 与现有工具的对比表；(3) 跨负样本策略的稳健性表。"
       "按期刊三线表规范排版，数值保留 3 位小数，最优值加粗。同时输出 LaTeX 和 Word 两种格式。"
       "参照 skills/00.1 的 references/08-tables-plots.md 规范。")

doc.add_heading("6.3 可解释性（让论文从「能预测」变成「有洞见」）", level=2)
b("SHAP：对 L1/L2 模型算特征重要性，看净电荷、疏水性、特定二肽是否如生物学预期般重要。")
b("注意力可视化：微调 ESM-2 时提取注意力权重，定位模型关注的残基位置。")
b("Motif 分析：把预测为高活性的序列做基序富集，与已知抗菌机制（如两亲性 α 螺旋）对照。")
b("反事实：对某条肽做单点突变扫描，看预测值如何变化，找出关键残基——这一步很容易做出漂亮的图。")

prompt("对最优模型做 SHAP 分析，输出全局特征重要性图和前 20 个重要特征的依赖图。"
       "然后对测试集中预测概率最高的 10 条肽做单点丙氨酸扫描，生成热图展示每个位置突变后预测值的变化。"
       "结合抗菌肽的膜作用机制解释这些发现，标注哪些结论有文献支持（用 paper-search 核实）、哪些是本研究的新观察。")

doc.add_page_break()

# ============ 7 写作与复现 ============
doc.add_heading("第 7 章　论文写作与复现包", level=1)

doc.add_heading("7.1 用 Skills 驱动写作", level=2)
table(
    ["章节", "调用的 Skill", "配套 MCP"],
    [
        ["Introduction", "03-K-Dense scientific-writing、36 综述技能", "paper-search + Zotero"],
        ["Methods", "00.1-Python 的 05-modeling / 06-robustness", "memory（取回每次决策记录）"],
        ["Results", "00.1 的 08-tables-plots", "jupyter + filesystem"],
        ["Discussion", "03-K-Dense hypothesis-generation", "paper-search（对照已有结论）"],
        ["数据可得性 / 复现", "54-open-science-skills", "github + filesystem"],
    ],
    widths=[1.4, 3.0, 2.6],
)

prompt("按 skills/54-open-science-skills 的规范整理复现包："
       "生成 environment.yml（锁定版本）、run_all.sh（从原始数据到所有图表的一键复现）、"
       "数据可得性声明（说明哪些数据可公开、哪些受限）、以及一份 CHANGELOG（从 memory 的研究日志汇总）。"
       "然后建 GitHub 仓库并打 v1.0 tag。")

doc.add_heading("7.2 投稿前自查清单", level=2)
table(
    ["项目", "自查"],
    [
        ["负样本构造规则是否完整交代", "关键词、长度、定位、CD-HIT 阈值缺一不可"],
        ["是否做了同源性感知划分", "随机划分会被审稿人直接质疑"],
        ["是否有独立外部测试集", "最好来自不同时间或不同数据库"],
        ["是否报告 MCC 与 SPE", "只报 Accuracy/F1 会被要求补充"],
        ["是否多种子重复", "单次结果不可接受"],
        ["与现有工具比较是否公平", "必须在同一测试集上重跑，不能引用原文数字"],
        ["代码与数据是否可获取", "多数期刊已强制要求"],
        ["AI 使用是否披露", "从 memory 导出工具使用日志"],
        ["每条参考文献是否真实存在", "用 Zotero 逐条核实 DOI"],
    ],
    widths=[3.2, 3.8],
)

doc.add_page_break()

# ============ 8 计划 ============
doc.add_heading("第 8 章　四周执行计划", level=1)
table(
    ["周次", "任务", "交付物", "主用工具"],
    [
        ["第 1 周", "环境配置 + 文献梳理 + 确定任务类型与研究问题", "文献对比表、研究设计文档", "MCP: paper-search/Zotero；Skill: 36/52"],
        ["第 2 周", "数据下载、清洗、负样本构造、同源性感知划分", "processed 数据集 + 样本流失表 + 划分索引", "jupyter、filesystem、DBHub"],
        ["第 3 周", "L0→L3 四级建模 + Optuna 调参 + 多种子重复", "metrics JSON、消融表", "Optuna、jupyter、context7"],
        ["第 4 周", "外部对比、稳健性检验、SHAP 可解释性、出图出表", "全部图表 + 结果章节草稿", "jupyter、Skill 00.1/08"],
        ["第 5 周起", "写作、复现包、投稿", "论文 + GitHub 仓库", "Zotero、github、Skill 54"],
    ],
    widths=[0.8, 2.4, 2.1, 1.7],
)

doc.add_heading("8.1 反重力的并行用法", level=2)
p("反重力的 Agent Manager 可以同时跑多个任务，建议这样分配（也是安全上的最佳实践）：", after=4)
b("Agent A（只读类工具）：跑文献检索与整理，只挂 paper-search、Zotero、markitdown。")
b("Agent B（计算类）：跑训练与调参，只挂 jupyter、optuna、context7。")
b("Agent C（文件类）：整理结果、写文档，只挂 filesystem、memory。")
p("不要让同一个 Agent 同时拥有「抓取外部内容」和「写本地文件/数据库」的能力——"
  "网页或 PDF 里可能藏有针对 AI 的注入指令，这是目前 MCP 生态最现实的安全风险。", size=10)

doc.add_heading("附录 A　常见错误速查", level=1)
table(
    ["错误", "后果", "正确做法"],
    [
        ["随机划分训练测试集", "指标虚高 5–15 个百分点", "CD-HIT 聚类后按簇划分"],
        ["合并多库后未去重", "同一条肽横跨训练与测试", "序列级完全去重 + 同源去冗余"],
        ["在全量数据上做标准化/特征选择", "隐性泄漏", "放进 Pipeline，只在训练折上 fit"],
        ["只报 Accuracy 和 F1", "掩盖低特异度问题", "主报 MCC，并给出 SEN/SPE"],
        ["引用他人论文里的指标做对比", "不公平比较，审稿人必挑", "在自己的测试集上重跑对方模型"],
        ["单次运行的结果", "不可复现", "≥5 个随机种子，报均值±标准差"],
        ["反复用测试集调策略", "测试集失去意义", "冻结测试集，只在最后用一次"],
        ["让 AI 生成参考文献", "可能凭空捏造", "只引 Zotero 库中真实条目，逐条核 DOI"],
        ["直接采信 AI 报告的数字", "结果错误", "所有进论文的数字手动重跑脚本确认"],
    ],
    widths=[2.0, 2.0, 3.0],
)

doc.add_heading("附录 B　关键参考方向", level=1)
b("负样本偏差与基准公平性：Sidorczuk 等, Briefings in Bioinformatics, 2022 —— 本课题设计的必读起点。")
b("蛋白语言模型用于 AMP 检测：近年多篇工作（ESM-2 / ProtT5 / ProtBERT 路线），可用 paper-search 拉取最新。")
b("菌种特异性与 MIC 回归：DBAASP v3、GRAMPA、ESKAPEE 类数据集相关工作。")
b("主流数据库：DBAASP v3、dbAMP、DRAMP 3.0、APD3、CAMPR4、LAMP2、YADAMP。")
p("注意：本文档中的数据库规模、指标区间等为写作时的概览性描述，具体数字请以你实际下载的版本与官方文档为准；"
  "文献结论请用 paper-search + Zotero 自行核实后再写进论文。",
  size=9.5, color=RGBColor(0x55, 0x55, 0x55))

doc.save(OUT)
print("saved:", OUT)

# -*- coding: utf-8 -*-
"""生成《Antigravity IDE 科研与办公 MCP 适配指南》docx。"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "Antigravity科研办公MCP适配指南.docx"
doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.25
for i, sz in [(1, 20), (2, 15), (3, 12.5)]:
    s = doc.styles[f"Heading {i}"]
    s.font.name = "Calibri"
    s.font.size = Pt(sz)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def p(text="", bold=False, italic=False, size=10.5, color=None, align=None, space_after=6):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        r.font.color.rgb = color
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(space_after)
    return par


def bullet(text):
    par = doc.add_paragraph(style="List Bullet")
    r = par.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    par.paragraph_format.space_after = Pt(2)
    return par


def code(text):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.2)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    par.paragraph_format.line_spacing = 1.0
    r = par.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F4F7")
    par._p.get_or_add_pPr().append(shd)
    return par


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------- 封面 ----------------
p("Antigravity IDE 科研 / 办公 MCP 适配指南", bold=True, size=23,
  color=RGBColor(0x1F, 0x3B, 0x63), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
p("哪些 MCP 在 Google Antigravity（反重力）里真能用 · 配置怎么改 · 复制即用", size=11.5,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))
p("配套文件：科研MCP服务推荐与安装指南.docx（VS Code 版） · examples/mcp-antigravity/", size=9.5,
  align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x77, 0x77, 0x77))
p("生成日期：2026-08-20", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER,
  color=RGBColor(0x77, 0x77, 0x77), space_after=14)

# ---------------- 1 结论先行 ----------------
doc.add_heading("1. 一句话结论", level=1)
p("Antigravity 完整支持 MCP 协议，所以上一份文档里推荐的服务器基本全部可用 —— 只是配置文件的位置和写法不同。"
  "真正的差异只有四条：", bold=True)
bullet("① 顶层键是 mcpServers（VS Code 是 servers）。")
bullet("② 远程服务器用 serverUrl 字段（VS Code 用 type + url）。")
bullet("③ 没有 ${input:...} 交互式密钥提示，改用 ${ENV_VAR} 环境变量占位或直接写 env。")
bullet("④ 只有一个例外真的不能用：GitHub Copilot 的远程端点 https://api.githubcopilot.com/mcp —— 它绑定 Copilot 订阅认证，"
       "在 Antigravity 里要改用 GitHub 官方本地 Docker 版 + PAT，或直接用 MCP Store 里的 GitHub 条目。")

p("Antigravity 有三种加装方式：", bold=True, space_after=2)
bullet("MCP Store（“…”菜单 → MCP Servers）：官方精选、点 Install 即可，但条目有限（GitHub / Figma / Linear / Notion / Supabase / MongoDB / Stripe / Sequential Thinking 等）。")
bullet("手工改 mcp_config.json（Manage MCP Servers → View raw config）：科研类服务器几乎都得走这条路。")
bullet("插件（Extensions）自带 MCP：装完插件后回 Manage MCP 点 Refresh 即可看到。")

doc.add_heading("1.1 配置文件在哪", level=2)
table(
    ["场景", "路径"],
    [
        ["Antigravity 2.0（全局，三端共用）", "~/.gemini/config/mcp_config.json"],
        ["Antigravity IDE 早期版本（全局）", "~/.gemini/antigravity/mcp_config.json"],
        ["工作区级（跟项目走，可进 Git）", "<项目根>/.agents/mcp_config.json"],
        ["Windows 用户目录写法", "C:\\Users\\<你>\\.gemini\\config\\mcp_config.json"],
        ["Antigravity CLI / SDK", "同上，全局 + .agents/mcp_config.json 自动发现"],
    ],
    widths=[2.6, 4.4],
)
p("改完文件回到 Manage MCP Servers 点 Refresh，几秒后新工具就会出现；不用重启 IDE。", size=10, italic=True)

doc.add_page_break()

# ---------------- 2 兼容矩阵 ----------------
doc.add_heading("2. 兼容性筛选表：哪些能用、要不要改", level=1)
p("✅ = 直接可用；🔧 = 可用但要改写法/换方案；⚠️ = 有前提条件；❌ = 不建议在 Antigravity 里用。", size=9.5,
  color=RGBColor(0x66, 0x66, 0x66))

doc.add_heading("2.1 科研向", level=2)
table(
    ["服务器", "Antigravity", "说明 / 改动点"],
    [
        ["Context7（最新库文档）", "✅", "stdio 版 npx @upstash/context7-mcp 最稳；远程版写 serverUrl。首推常开"],
        ["Markitdown（PDF/Office→MD）", "✅", "uvx markitdown-mcp，零改动"],
        ["paper-search-mcp（20+ 文献库）", "✅", "uvx --from git+… 直接跑，注意关掉 Sci-Hub 通道"],
        ["Zotero MCP", "✅", "本地模式 ZOTERO_LOCAL=true，零改动"],
        ["OpenAlex / arXiv 类", "✅", "均为 stdio，零改动"],
        ["SearXNG Search", "✅", "需自建或指定实例 URL"],
        ["Tavily", "✅", "env 里写 TAVILY_API_KEY，不支持 input 提示框"],
        ["Firecrawl", "✅", "同上；付费，注意配额"],
        ["DBHub（SQL 数据库）", "✅", "务必带 --readonly"],
        ["mcp-stata / Jupyter MCP", "⚠️", "可用；需本机 Stata 授权 / 已启动的 Jupyter 服务"],
        ["Playwright", "⚠️", "可用，但 Antigravity 自带 Browser Use（浏览器代理）功能，多数情况不必再装"],
        ["Serena（语义代码检索）", "🔧", "可用；但 Antigravity 自身代码索引已较强，装了会与内置工具重叠"],
        ["Basic Memory（本地知识库）", "✅", "推荐；与 Antigravity 的 Knowledge/Artifacts 互补"],
        ["GitHub（Copilot 远程端点）", "❌", "改用 Store 里的 GitHub，或本地 Docker 版 + PAT（见 4.3）"],
        ["Desktop Commander（终端）", "❌", "不必装。Antigravity 的 Agent 本身就有终端与文件权限，叠加只会放大风险"],
    ],
    widths=[1.9, 0.9, 4.2],
)

doc.add_heading("2.2 办公向", level=2)
table(
    ["服务器", "Antigravity", "能干什么（办公场景）"],
    [
        ["Office-Word MCP（GongRzhe）", "✅", "生成/编辑 .docx：改标题样式、批量替换、加表格与页眉、合并文档、加批注 —— 写报告、周报、结题材料"],
        ["Excel MCP（haris-musa）", "✅", "读写 .xlsx：公式、透视、图表、多表合并、批量格式化 —— 处理台账、经费表、问卷原始表"],
        ["Markitdown", "✅", "反向操作：把收到的 PDF/PPT/扫描件/会议录音统一转成文本再处理"],
        ["Google Workspace MCP（workspace-mcp）", "⚠️", "Gmail / Drive / Docs / Sheets / Slides / Calendar / Tasks 全套；需自建 Google Cloud OAuth 凭据"],
        ["Notion（Store 内有）", "✅", "点 Install 即可，读写研究看板、任务表"],
        ["Filesystem MCP", "✅", "批量整理下载目录、按规则归档文件（限定可访问目录）"],
        ["Sequential Thinking（Store 内有）", "✅", "结构化拆解复杂任务，写长报告/答辩逻辑时有帮助"],
        ["Time / Fetch 等官方小工具", "✅", "时区换算、抓单个网页，轻量补齐"],
    ],
    widths=[2.1, 0.9, 4.0],
)

doc.add_page_break()

# ---------------- 3 复制即用配置 ----------------
doc.add_heading("3. 复制即用的 mcp_config.json", level=1)
p("粘贴到 ~/.gemini/config/mcp_config.json（全局）或项目里的 .agents/mcp_config.json（推荐，随项目走）。"
  "建议一次只开 5–8 个服务器。", size=10, italic=True)

doc.add_heading("3.1 科研套装（文献 → 精读 → 记录）", level=2)
code('''{
  "mcpServers": {
    "context7": {
      "command": "npx",
      "args": ["-y", "@upstash/context7-mcp"]
    },
    "markitdown": {
      "command": "uvx",
      "args": ["markitdown-mcp"]
    },
    "paper-search": {
      "command": "uvx",
      "args": ["--from", "git+https://github.com/openags/paper-search-mcp", "paper-search-mcp"],
      "env": { "PAPER_SEARCH_MCP_UNPAYWALL_EMAIL": "your@email.edu" }
    },
    "zotero": {
      "command": "uvx",
      "args": ["--from", "git+https://github.com/54yyyu/zotero-mcp", "zotero-mcp"],
      "env": { "ZOTERO_LOCAL": "true" }
    },
    "memory": {
      "command": "uvx",
      "args": ["basic-memory", "mcp"],
      "env": { "BASIC_MEMORY_HOME": "/Users/你的用户名/research-notes" }
    }
  }
}''')

doc.add_heading("3.2 数据分析套装", level=2)
code('''{
  "mcpServers": {
    "dbhub": {
      "command": "npx",
      "args": ["-y", "@bytebase/dbhub", "--transport", "stdio",
               "--dsn", "sqlite:////Users/你的用户名/project/data/panel.db",
               "--readonly"]
    },
    "stata": {
      "command": "uvx",
      "args": ["mcp-stata"],
      "env": { "STATA_PATH": "/Applications/Stata/StataMP.app/Contents/MacOS/StataMP" }
    },
    "excel": {
      "command": "uvx",
      "args": ["excel-mcp-server", "stdio"],
      "env": { "EXCEL_FILES_PATH": "/Users/你的用户名/project/data" }
    },
    "context7": { "command": "npx", "args": ["-y", "@upstash/context7-mcp"] }
  }
}''')

doc.add_heading("3.3 办公套装（Word / Excel / Google Workspace）", level=2)
code('''{
  "mcpServers": {
    "word": {
      "command": "uvx",
      "args": ["--from", "office-word-mcp-server", "word_mcp_server"]
    },
    "excel": {
      "command": "uvx",
      "args": ["excel-mcp-server", "stdio"],
      "env": { "EXCEL_FILES_PATH": "/Users/你的用户名/Documents/work" }
    },
    "markitdown": { "command": "uvx", "args": ["markitdown-mcp"] },
    "google-workspace": {
      "command": "uvx",
      "args": ["workspace-mcp", "--tool-tier", "core"],
      "env": {
        "GOOGLE_OAUTH_CLIENT_ID": "${GOOGLE_OAUTH_CLIENT_ID}",
        "GOOGLE_OAUTH_CLIENT_SECRET": "${GOOGLE_OAUTH_CLIENT_SECRET}",
        "USER_GOOGLE_EMAIL": "you@example.com",
        "OAUTHLIB_INSECURE_TRANSPORT": "1"
      }
    },
    "filesystem": {
      "command": "npx",
      "args": ["-y", "@modelcontextprotocol/server-filesystem",
               "/Users/你的用户名/Documents/work"]
    }
  }
}''')
p("说明：Google Workspace 需先在 Google Cloud Console 建一个 OAuth 桌面应用凭据并启用 Gmail/Drive/Docs/Sheets API；"
  "只想要“别人发我的 docx/xlsx 我来批量处理”，用 word + excel + markitdown 三个本地服务器就够了，零配置零成本。", size=10)

doc.add_heading("3.4 远程服务器写法（与 VS Code 不同，注意）", level=2)
code('''{
  "mcpServers": {
    "some-remote": {
      "serverUrl": "https://api.example.com/mcp/",
      "headers": { "Authorization": "Bearer ${MY_TOKEN}" }
    },
    "oauth-server": {
      "serverUrl": "https://api.example.com/mcp/"
    },
    "gcp-service": {
      "serverUrl": "https://example.googleapis.com/mcp/",
      "authProviderType": "google_credentials"
    }
  }
}''')
p("支持动态客户端注册（DCR）的远程服务器不用填任何密钥，Antigravity 会自动走 OAuth 弹窗。"
  "若某远程服务只给了 SSE/HTTP 而 Antigravity 连不上，用 npx mcp-remote <url> 包一层转成 stdio 即可（社区通行做法）。", size=10)

doc.add_page_break()

# ---------------- 4 迁移与排错 ----------------
doc.add_heading("4. 从 VS Code 配置迁移过来：逐条对照", level=1)
table(
    ["VS Code (.vscode/mcp.json)", "Antigravity (mcp_config.json)"],
    [
        ['顶层 "servers": { … }', '顶层 "mcpServers": { … }'],
        ['"type": "stdio" + command/args', '省略 type，直接写 command/args'],
        ['"type": "http", "url": "…"', '"serverUrl": "…"'],
        ['"headers": {…}', '"headers": {…}（相同）'],
        ['"inputs" + ${input:key} 密钥提示', '不支持；改 env 直写或 ${ENV_VAR} 环境变量占位'],
        ['${workspaceFolder}', '不保证支持，建议写绝对路径'],
        ['"sandboxEnabled" / sandbox 配置', '无；改用 Antigravity 的权限策略 mcp(server/tool)'],
        ['扳手图标勾选启用工具', 'Manage MCP Servers 里逐个 toggle 开关'],
    ],
    widths=[3.3, 3.7],
)

doc.add_heading("4.1 权限控制（Antigravity 特有，务必设置）", level=2)
bullet("默认所有未配置的 MCP 工具处于 Ask 模式：每次调用都要你点同意。科研数据敏感，建议保持默认。")
bullet("要放行可在策略里写：mcp(zotero/*) 放行整个服务器、mcp(dbhub/query) 只放行单个工具、mcp(*) 全放行（不推荐）。")
bullet("写操作类工具（Excel 写入、Word 保存、数据库写、Drive 删除）永远别加进自动放行清单。")

doc.add_heading("4.2 常见故障（Antigravity 用户高频踩坑）", level=2)
table(
    ["现象", "原因", "解决"],
    [
        ['exec: "npx": executable file not found', "Antigravity 启动时未继承 shell 的 PATH", "把 command 换成绝对路径，如 /usr/local/bin/npx、~/.local/bin/uvx；或用 /bin/sh -c 包一层并显式设 PATH"],
        ["装完 Store 里的服务器一直失败", "本机没装 Node.js 或版本过低", "装 Node 22+ 并重启电脑（不是重启 IDE），再重新 Install"],
        ["改了 json 没生效", "没点刷新", "Manage MCP Servers → Refresh"],
        ["uvx 找不到包", "首次拉取超时/网络问题", "先在终端手动跑一次同样的 uvx 命令，确认能装上再回 IDE"],
        ["Windows 路径报错", "反斜杠未转义", '写 "C:\\\\Users\\\\me\\\\data" 或直接用正斜杠'],
        ["工具太多、Agent 乱调", "同时开了十几个服务器", "只开当前阶段用的；其余在 Manage 里 toggle 关掉"],
    ],
    widths=[2.2, 1.7, 3.1],
)

doc.add_heading("4.3 GitHub 在 Antigravity 里怎么接", level=2)
p("优先：MCP Store 里直接 Install GitHub（走 OAuth，最省事）。若 Store 版本受限，用官方本地 Docker 版：", space_after=4)
code('''{
  "mcpServers": {
    "github": {
      "command": "docker",
      "args": ["run", "-i", "--rm",
               "-e", "GITHUB_PERSONAL_ACCESS_TOKEN",
               "ghcr.io/github/github-mcp-server"],
      "env": { "GITHUB_PERSONAL_ACCESS_TOKEN": "${GITHUB_PAT}" }
    }
  }
}''')
p("PAT 只勾 repo 必要权限；未发表的数据与代码请确认仓库为 private 后再让 Agent 操作。", size=10)

doc.add_page_break()

# ---------------- 5 落地建议 ----------------
doc.add_heading("5. 给你的落地建议（按优先级）", level=1)
table(
    ["优先级", "装什么", "为什么", "成本"],
    [
        ["★★★ 第一批", "Context7 + Markitdown + Zotero", "零成本零风险，立刻解决“读 PDF / 不编造引用 / 包 API 用错”三个高频痛点", "免费"],
        ["★★★ 第一批", "Word MCP + Excel MCP", "办公提效最直接：报告排版、表格批处理，纯本地不联网", "免费"],
        ["★★ 第二批", "paper-search + Basic Memory", "打通“检索→精读→综述矩阵→留痕”的完整链路", "免费"],
        ["★★ 第二批", "DBHub（只读）或 mcp-stata", "把实证分析环节接进 Agent，实现跑完即读结果", "免费（需本机 Stata）"],
        ["★ 按需", "Google Workspace / Notion", "多人协作、日程与邮件汇总；配置成本较高", "免费但要建 OAuth"],
        ["★ 按需", "Tavily / Firecrawl", "需要外部网页数据时再开，先设预算上限", "付费"],
        ["不建议", "Desktop Commander / Serena", "与 Antigravity 内置的终端、代码索引、浏览器能力重叠，徒增风险和工具噪音", "—"],
    ],
    widths=[1.0, 1.9, 3.1, 1.0],
)

doc.add_heading("5.1 三条 Antigravity 特有的用法提醒", level=2)
bullet("Antigravity 是 Agent-first 的（Agent Manager 可并行跑多个任务）。建议把“检索类”与“写文件类”任务分给不同 Agent 会话，避免一个 Agent 同时握有抓取和写盘能力（提示注入风险）。")
bullet("它会产出 Artifacts（任务清单、walkthrough、截图）。把 Basic Memory 接上后，可让 Agent 把每次实证结果与决策写成 Markdown 留痕，投稿时直接用作 AI 使用披露材料。")
bullet("工作区级 .agents/mcp_config.json 建议提交进论文项目仓库，作为可复现研究环境的一部分 —— 合作者克隆后即拥有同样的工具链。")

doc.add_heading("附：信息来源与时效", level=2)
p("本文依据 Google Antigravity 官方 MCP 文档（antigravity.google/docs/mcp，含 Antigravity 2.0 / IDE / CLI / SDK 四种形态的配置说明）"
  "与社区实测记录整理，核对时间 2026-08-20。Antigravity 迭代较快，配置路径可能从 ~/.gemini/antigravity/ 迁移到 ~/.gemini/config/，"
  "以你本机 “View raw config” 实际打开的文件为准。第三方服务器的安装命令以其 README 为准。", size=9.5,
  color=RGBColor(0x55, 0x55, 0x55))

doc.save(OUT)
print("saved:", OUT)
